"""[역할 3] export: clean_news + 기사별 최신 요약 → CSV·JSONL·XLSX.

- --status(all·cleaned·summarized·summary_failed·no_body), 카테고리, 기간으로 거른다.
- 기사 본문은 기본으로 빼고(저작권), --include-body일 때만 넣는다.
- CSV는 엑셀에서 한글이 깨지지 않게 BOM이 붙은 UTF-8(utf-8-sig)로 쓴다.
- 파일 이름: news_{상태}_{YYYYMMDD_HHMMSS}.{형식}
"""
import csv
import json
import logging
from datetime import datetime

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill

from news_cli.db import KST, now_str
from news_cli.output.common import select_news

log = logging.getLogger(__name__)

FORMATS = ("csv", "jsonl", "xlsx")
COLUMNS = ("id", "published_at", "category", "category_source", "section", "title", "url", "author", "methods",
           "status", "body_source", "summary", "keywords", "summary_model", "summarized_at",
           "first_collected_at", "last_collected_at")
XLSX_WIDTHS = {"published_at": 19, "category": 14, "title": 50, "url": 30, "summary": 80, "keywords": 30, "body": 80,
               "summarized_at": 19, "first_collected_at": 19, "last_collected_at": 19}
EXCEL_CELL_LIMIT = 32767


def parse_formats(value):
    """'csv,xlsx' → ['csv', 'xlsx']. 모르는 형식이 있으면 ValueError."""
    formats = list(dict.fromkeys(name.strip().lower() for name in value.split(",") if name.strip()))
    unknown = [name for name in formats if name not in FORMATS]
    if unknown or not formats:
        raise ValueError(f"알 수 없는 형식: {', '.join(unknown) or value} (사용 가능: {', '.join(FORMATS)})")
    return formats


def export_rows(conn, filters, include_body=False):
    """내보낼 행(dict) 목록. 열 순서는 COLUMNS(+ body). keywords는 목록 그대로 둔다."""
    columns = COLUMNS + (("body",) if include_body else ())
    return [{column: article.get(column) for column in columns} for article in select_news(conn, filters)]


def _flat(value):
    """CSV·XLSX 칸 값: 목록은 '; '로 잇고 None은 빈 칸."""
    if isinstance(value, list):
        return "; ".join(value)
    return "" if value is None else value


def write_csv(rows, path):
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows({key: _flat(value) for key, value in row.items()} for row in rows)


def write_jsonl(rows, path):
    with path.open("w", encoding="utf-8") as f:
        for row in rows:
            f.write(json.dumps(row, ensure_ascii=False) + "\n")


def _xlsx_value(value):
    value = _flat(value)
    if isinstance(value, str):
        value = ILLEGAL_CHARACTERS_RE.sub("", value)[:EXCEL_CELL_LIMIT]  # 엑셀이 받지 않는 제어 문자·길이
    return value


def write_xlsx(rows, path, info):
    wb = Workbook()
    ws = wb.active
    ws.title = "news"
    columns = list(rows[0])
    ws.append(columns)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDE7F3")
    for row in rows:
        ws.append([_xlsx_value(row[column]) for column in columns])
    ws.freeze_panes = "A2"
    for index, column in enumerate(columns, 1):
        letter = ws.cell(row=1, column=index).column_letter
        ws.column_dimensions[letter].width = XLSX_WIDTHS.get(column, 12)
        if column in ("summary", "body"):
            for cell in ws[letter][1:]:
                cell.alignment = Alignment(wrap_text=True, vertical="top")

    sheet = wb.create_sheet("info")
    for key, value in info.items():
        sheet.append([key, value])
    sheet.column_dimensions["A"].width = 14
    sheet.column_dimensions["B"].width = 60
    wb.save(path)


def run_export(conn, filters, formats, output_dir, include_body=False):
    """조건에 맞는 기사를 형식별 파일로 쓰고 [(형식, 경로)]를 돌려준다. 기사가 없으면 빈 목록(파일을 만들지 않음)."""
    rows = export_rows(conn, filters, include_body)
    if not rows:
        log.warning("조건에 맞는 기사가 없어 내보내지 않습니다 (%s)", filters.describe())
        return []
    output_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now(KST).strftime("%Y%m%d_%H%M%S")
    status = filters.status or "all"
    info = {"생성 시각": now_str(), "조건": filters.describe(), "기사 수": len(rows),
            "본문 포함": "예" if include_body else "아니오 (--include-body로 포함)"}
    written = []
    for fmt in formats:
        path = output_dir / f"news_{status}_{stamp}.{fmt}"
        if fmt == "csv":
            write_csv(rows, path)
        elif fmt == "jsonl":
            write_jsonl(rows, path)
        else:
            write_xlsx(rows, path, info)
        written.append((fmt, path))
        log.info("내보내기: %s %d건 → %s", fmt.upper(), len(rows), path)
    if not include_body:
        log.info("기사 본문은 넣지 않았습니다 (필요하면 --include-body, 저작권상 외부 공개 금지)")
    return written
