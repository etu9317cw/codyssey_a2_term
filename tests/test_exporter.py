"""내보내기: --status 필터, 기본은 본문 제외, CSV·JSONL·XLSX 생성 후 다시 읽어 확인 (메모리 DB, 임시 폴더)."""
import csv
import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from news_cli import db
from news_cli.output import exporter
from news_cli.output.common import Filters
from tests.support import add_clean_news, add_summary


class ExporterTest(unittest.TestCase):
    def setUp(self):
        self.conn = db.connect(":memory:")
        self.done = add_clean_news(self.conn, "1", "HVDC 입찰\x01", category="전력", status="summarized",
                                   published_at="2026-10-09 09:00:00")
        add_summary(self.conn, self.done, "옛 요약", ["옛"])
        add_summary(self.conn, self.done, "HVDC 입찰방안을 공개했다.", ["HVDC", "한국전력"])
        add_clean_news(self.conn, "2", "수상태양광", category="태양광", published_at="2026-10-08 09:00:00")
        add_clean_news(self.conn, "3", "본문 없음", body="", category="태양광", status="no_body",
                       published_at="2026-10-07 09:00:00")
        self.tmp = tempfile.TemporaryDirectory()
        self.out = Path(self.tmp.name) / "exports"

    def tearDown(self):
        self.conn.close()
        self.tmp.cleanup()

    def test_status_filter(self):
        rows = exporter.export_rows(self.conn, Filters(status="summarized"))
        self.assertEqual([r["id"] for r in rows], [self.done])
        self.assertEqual(rows[0]["summary"], "HVDC 입찰방안을 공개했다.")  # 최신 요약
        self.assertEqual(rows[0]["keywords"], ["HVDC", "한국전력"])
        self.assertEqual(len(exporter.export_rows(self.conn, Filters(status="all"))), 3)
        self.assertEqual([r["id"] for r in exporter.export_rows(self.conn, Filters(categories=["태양광"],
                                                                                   date_from="2026-10-08"))], [2])

    def test_body_excluded_by_default(self):
        self.assertNotIn("body", exporter.export_rows(self.conn, Filters())[0])
        self.assertIn("body", exporter.export_rows(self.conn, Filters(), include_body=True)[0])

    def test_three_formats_written_and_readable(self):
        written = dict(exporter.run_export(self.conn, Filters(status="summarized"), ["csv", "jsonl", "xlsx"], self.out))
        self.assertEqual(set(written), {"csv", "jsonl", "xlsx"})
        self.assertTrue(written["csv"].name.startswith("news_summarized_"))

        with written["csv"].open(encoding="utf-8-sig", newline="") as f:
            rows = list(csv.DictReader(f))
        self.assertEqual(list(rows[0]), list(exporter.COLUMNS))
        self.assertEqual((len(rows), rows[0]["keywords"], rows[0]["status"]), (1, "HVDC; 한국전력", "summarized"))

        lines = written["jsonl"].read_text(encoding="utf-8").splitlines()
        self.assertEqual(len(lines), 1)
        self.assertEqual(json.loads(lines[0])["keywords"], ["HVDC", "한국전력"])

        wb = load_workbook(written["xlsx"])
        sheet = wb["news"]
        self.assertEqual([c.value for c in sheet[1]], list(exporter.COLUMNS))
        self.assertEqual(sheet.max_row, 2)
        self.assertEqual(sheet.cell(row=2, column=6).value, "HVDC 입찰")  # 엑셀이 못 받는 제어 문자 제거
        self.assertIn("info", wb.sheetnames)
        wb.close()

    def test_include_body_column(self):
        written = dict(exporter.run_export(self.conn, Filters(status="cleaned"), ["csv"], self.out, include_body=True))
        with written["csv"].open(encoding="utf-8-sig", newline="") as f:
            row = next(csv.DictReader(f))
        self.assertTrue(row["body"].startswith("본문 문장입니다."))

    def test_no_rows_writes_nothing(self):
        self.assertEqual(exporter.run_export(self.conn, Filters(categories=["수소"]), ["csv"], self.out), [])
        self.assertFalse(self.out.exists())

    def test_parse_formats(self):
        self.assertEqual(exporter.parse_formats("CSV, xlsx,csv"), ["csv", "xlsx"])
        for bad in ("pdf", "csv,txt", ","):
            with self.assertRaises(ValueError):
                exporter.parse_formats(bad)


if __name__ == "__main__":
    unittest.main()
